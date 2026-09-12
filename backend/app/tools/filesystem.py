"""
Filesystem Tool — Full file & folder CRUD + File Generation
Supports: PDF, Excel (xlsx), CSV, DOCX, TXT, JSON, Markdown, Python, and any text extension.
"""
import os
import json
import shutil
from pathlib import Path
from typing import Optional, List, Dict, Any

# File generation libraries
from fpdf import FPDF
import openpyxl
import pandas as pd
from docx import Document


# ── File & Folder CRUD ────────────────────────────────────────────────────────

def list_directory(path: str) -> Dict:
    """List all files and folders in a directory."""
    p = Path(path)
    if not p.exists():
        return {"error": f"Path does not exist: {path}"}
    items = [{"name": item.name, "type": "dir" if item.is_dir() else "file", "size": item.stat().st_size if item.is_file() else None}
             for item in sorted(p.iterdir())]
    return {"path": str(p.resolve()), "items": items}


def read_file(path: str) -> str:
    """Read a text file and return its contents."""
    try:
        return Path(path).read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"Error reading file: {e}"


def write_file(path: str, content: str) -> str:
    """Write content to a file, creating parent directories as needed."""
    try:
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
        return f"✅ File written: {p.resolve()}"
    except Exception as e:
        return f"Error writing file: {e}"


def delete_file_or_folder(path: str) -> str:
    """Delete a file or folder (recursively)."""
    try:
        p = Path(path)
        if p.is_file():
            p.unlink()
        elif p.is_dir():
            shutil.rmtree(p)
        return f"✅ Deleted: {path}"
    except Exception as e:
        return f"Error deleting: {e}"


def create_folder(path: str) -> str:
    """Create a folder and all intermediate directories."""
    try:
        Path(path).mkdir(parents=True, exist_ok=True)
        return f"✅ Folder created: {path}"
    except Exception as e:
        return f"Error creating folder: {e}"


def copy_item(src: str, dest: str) -> str:
    """Copy a file or folder."""
    try:
        s, d = Path(src), Path(dest)
        if s.is_file():
            shutil.copy2(s, d)
        else:
            shutil.copytree(s, d)
        return f"✅ Copied {src} → {dest}"
    except Exception as e:
        return f"Error copying: {e}"


def move_item(src: str, dest: str) -> str:
    """Move/rename a file or folder."""
    try:
        shutil.move(src, dest)
        return f"✅ Moved {src} → {dest}"
    except Exception as e:
        return f"Error moving: {e}"


def search_files(root: str, pattern: str) -> List[str]:
    """Search for files matching a glob pattern."""
    return [str(p) for p in Path(root).rglob(pattern)]


# ── PDF Generation ─────────────────────────────────────────────────────────────

def create_pdf(path: str, title: str, content: str) -> str:
    """Generate a PDF file with a title and multi-line content."""
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, title, ln=True, align="C")
        pdf.ln(5)
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 8, content)
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(p))
        return f"✅ PDF created: {p.resolve()}"
    except Exception as e:
        return f"Error creating PDF: {e}"


# ── Excel Generation ───────────────────────────────────────────────────────────

def create_excel(path: str, sheet_name: str, data: List[Dict]) -> str:
    """Generate an Excel (.xlsx) file from a list of dicts."""
    try:
        df = pd.DataFrame(data)
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(str(p), engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        return f"✅ Excel created: {p.resolve()}"
    except Exception as e:
        return f"Error creating Excel: {e}"


# ── CSV Generation ─────────────────────────────────────────────────────────────

def create_csv(path: str, data: List[Dict]) -> str:
    """Generate a CSV file from a list of dicts."""
    try:
        df = pd.DataFrame(data)
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        df.to_csv(str(p), index=False)
        return f"✅ CSV created: {p.resolve()}"
    except Exception as e:
        return f"Error creating CSV: {e}"


# ── DOCX Generation ────────────────────────────────────────────────────────────

def create_docx(path: str, title: str, paragraphs: List[str]) -> str:
    """Generate a Word (.docx) document."""
    try:
        doc = Document()
        doc.add_heading(title, level=1)
        for para in paragraphs:
            doc.add_paragraph(para)
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(p))
        return f"✅ DOCX created: {p.resolve()}"
    except Exception as e:
        return f"Error creating DOCX: {e}"


# ── JSON Generation ────────────────────────────────────────────────────────────

def create_json_file(path: str, data: Any) -> str:
    """Write a JSON file."""
    try:
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        return f"✅ JSON created: {p.resolve()}"
    except Exception as e:
        return f"Error creating JSON: {e}"

